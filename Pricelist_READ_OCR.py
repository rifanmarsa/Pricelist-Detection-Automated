# 1. Connect Google Colab to Google Drive
from google.colab import drive
drive.mount('/content/drive')

import os
import requests
import datetime
import re
import cv2
import numpy as np
import pytesseract
import io
import pandas as pd
from PIL import Image as PILImage
from openpyxl import load_workbook

# --- DATA WRANGLING ---
# 1. Read Excel dataset file
input_excel_file = 'dummy_survey_report.xlsx'
df = pd.read_excel(input_excel_file)

# 2. Store static column references (first 24 columns)
static_cols = df.columns[:24]
melted_dfs = []

# 3. Melting Process: Combine 5 iterations of column groups
for i in range(5):
    start_col = 24 + (i * 3)
    if start_col + 2 < len(df.columns):
        temp_df = df[static_cols].copy()
        temp_df['Competitor Lubricant Brand'] = df.iloc[:, start_col]
        temp_df['Pricelist Effective Date'] = df.iloc[:, start_col + 1]
        temp_df['Pricelist Photo'] = df.iloc[:, start_col + 2]
        melted_dfs.append(temp_df)

final_df = pd.concat(melted_dfs, ignore_index=True)

# 4. Filtering: Keep rows containing valid URLs
final_df = final_df[final_df['Pricelist Photo'].astype(str).str.contains('http', na=False)]

# 5. Unnesting Data (Explode Function)
# Convert comma-separated URL strings into lists
final_df['Pricelist Photo'] = final_df['Pricelist Photo'].astype(str).str.split(',')
# Explode: Break list contents into independent new rows
final_df = final_df.explode('Pricelist Photo', ignore_index=True)
# Clean up excess whitespace characters
final_df['Pricelist Photo'] = final_df['Pricelist Photo'].str.strip()

# 6. Save data processing results to Excel format
output_file = 'final_survey_exploded_result.xlsx'
final_df.to_excel(output_file, index=False)
print(f"Done! Found {final_df.shape[0]} link rows. Saved as '{output_file}'")


# --- COMPUTER VISION & DATA EXTRACTION ---
def has_keywords(image_bytes):
    """
    Extracts text from images using OCR and validates pricing patterns.
    This function detects generic patterns (e.g., '000' number chains or 'Rp' combinations)
    to distinguish genuine pricelists from promotional banners.
    """
    try:
        # 1. Image Pre-processing
        np_arr = np.frombuffer(image_bytes, np.uint8)
        img_cv = cv2.imdecode(np_arr, cv2.IMREAD_COLOR)
        if img_cv is None: return False

        gray = cv2.cvtColor(img_cv, cv2.COLOR_BGR2GRAY)

        if np.mean(gray) < 5.0:
            print("      [-] Discarded: Image detected as black/blank")
            return False

        # Interpolation and Thresholding to clarify small text
        gray = cv2.resize(gray, None, fx=2, fy=2, interpolation=cv2.INTER_CUBIC)
        _, thresh = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
        img_pil = PILImage.fromarray(thresh)
        
        # 2. Text Extraction (OCR)
        # Whitelist specific characters to reduce OCR noise
        custom_config = r'--psm 11 -c tessedit_char_whitelist="abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789., "'
        text = pytesseract.image_to_string(img_pil, config=custom_config).lower()

        clean_text = " ".join(text.split())
        print(f"      [OCR] Text read: {clean_text[:120]}...")

        # 3. Pricelist Authenticity Filter Logic
        
        # Check explicit keywords
        explicit_kws = ["daftar harga", "pricelist", "price list",
                        "harga pelumas", "harga oli", "produk",
                        "product", "harga", "pcmo",
                        "mco", "hddo", "packaging",
                        "drum", "pail", "oil",
                        "oils", "lube", "lubes",
                        "lubricant", "lubricants", "item",
                        "items", "gear", "gardan",
                        "radiator", "coolant", "eceran"]
        if any(kw in text for kw in explicit_kws):
            print("      [!] Passed: Main keyword found")
            return True

        # Check price sequence patterns (main validation key)
        count_000 = text.count("000")
        if count_000 >= 3:
            print(f"      [!] Passed: Detected {count_000} price sequences")
            return True

        # Check currency and lubricant specification combinations
        tech_kws = ["10w", "20w", "15w", "5w", "800ml"]
        found_tech = sum(1 for kw in tech_kws if kw in text)
        if "rp" in text and found_tech >= 1:
            print("      [!] Passed: Combination of 'Rp' and viscosity/volume specifications found")
            return True

        return False

    except Exception as e:
        print(f"      [-] OCR Error: {e}")
        return False


def download_images_to_drive(input_file, sheet_name, url_col, outlet_col, frontliner_col, date_col, drive_folder):
    """
    Downloads images from URLs within Excel, validates them via OCR,
    and saves them to Google Drive with a structured naming format.
    """
    print(f"\nReading file {input_file}...")
    wb = load_workbook(input_file, data_only=True)

    if sheet_name not in wb.sheetnames:
        print(f"[-] Error: Sheet '{sheet_name}' not found!")
        return

    ws = wb[sheet_name]
    print(f"Processing sheet: '{sheet_name}'...")

    if not os.path.exists(drive_folder):
        os.makedirs(drive_folder)
        print(f"Folder '{drive_folder}' successfully created in Google Drive.")

    success = 0
    failed = 0
    skipped = 0

    for row in range(2, ws.max_row + 1):
        if ws.row_dimensions[row].hidden:
            continue

        url = ws[f"{url_col}{row}"].value

        if url and isinstance(url, str) and url.startswith('http'):
            try:
                # 1. Data Fetching
                outlet_id = ws[f"{outlet_col}{row}"].value
                frontliner_id = ws[f"{frontliner_col}{row}"].value
                date_val = ws[f"{date_col}{row}"].value

                # 2. File Naming Sanitation
                outlet_str = str(outlet_id) if outlet_id is not None else f"UnknownOutlet_Row{row}"
                outlet_str = re.sub(r'[\\/*?:"<>|]', "-", outlet_str)

                frontliner_str = str(frontliner_id) if frontliner_id is not None else f"UnknownFL_Row{row}"
                frontliner_str = re.sub(r'[\\/*?:"<>|]', "-", frontliner_str)

                # 3. Time Formatting
                if isinstance(date_val, datetime.datetime):
                    date_str = date_val.strftime("%Y-%m-%d_%H-%M-%S")
                else:
                    date_str = str(date_val) if date_val is not None else "UnknownDate"
                    date_str = re.sub(r'[\\/*?:"<>|]', "-", date_str)

                image_filename = f"{outlet_str}_{frontliner_str}_{date_str}_Row{row}.jpg"
                file_path = os.path.join(drive_folder, image_filename)

                print(f"Processing: {image_filename} (Row {row})")

                # 4. Downloading and Validation
                response = requests.get(url, stream=True, timeout=10)
                if response.status_code == 200:
                    image_data = response.content

                    if has_keywords(image_data):
                        with open(file_path, 'wb') as f:
                            f.write(image_data)
                        success += 1
                        print(f"   [+] SAVED")
                    else:
                        skipped += 1
                        print(f"   [-] DISCARDED (Not a Pricelist Pattern)")
                else:
                    failed += 1
                    print(f"   [-] Failed to download row {row}, Status Code: {response.status_code}")

            except Exception as e:
                failed += 1
                print(f"   [-] Error in row {row}: {e}")

    print(f"\n--- PROCESS COMPLETED ---")
    print(f"Successfully saved: {success} images")
    print(f"Not a pricelist   : {skipped} images (discarded)")
    print(f"Errors/Failed dl  : {failed} images")

if __name__ == "__main__":
    SOURCE_FILE = "final_survey_exploded_result.xlsx"
    SHEET_NAME = "Sheet1"
    
    URL_COL = "AA"
    OUTLET_COL = "J"
    FRONTLINER_COL = "F"
    DATE_COL = "B"
    
    TARGET_FOLDER = "/content/drive/MyDrive/Pricelist_Photos_TEST2"

    download_images_to_drive(SOURCE_FILE, SHEET_NAME, URL_COL, OUTLET_COL, FRONTLINER_COL, DATE_COL, TARGET_FOLDER)